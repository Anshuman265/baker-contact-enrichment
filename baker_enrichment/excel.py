"""Workbook validation, Baker record import, and deterministic Excel exports."""

from __future__ import annotations

from collections import Counter
from collections.abc import Iterable, Sequence
from dataclasses import asdict
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from baker_enrichment.api.contracts import ReverseLookupData
from baker_enrichment.models import BakerRecord, EnrichedContact, FailedRecord, ReverseLookupRow, ReviewRecord, RunReport


ENRICHED_COLUMNS = (
    "Lead ID", "First Name", "Last Name", "Email Domain", "LinkedIn Profile", "Headline", "Current Title",
    "Current Company", "Skill", "City", "State", "Country", "Company Domain", "Company LinkedIn",
    "Employee Count", "SIC Industry Classification", "NAICS Industry Classification", "Revenue", "Specialties",
    "Founded Year", "HQ City", "HQ State", "HQ Country", "Last Funding Type", "Last Funding Date", "Email", "Phone",
)
REVERSE_LOOKUP_COLUMNS = (
    "ID", "Display Name", "First Name", "Last Name", "Headline", "Current Role", "Current Company", "Summary",
    "Company", "Location", "Photo URL", "Profile URL", "Report Profile URL", "Connection Count",
    "Connections Obfuscated", "Public Profile", "Lookup Email", "Skills", "Locale Country", "Locale Language",
    "Education", "Work History",
)
REVIEW_COLUMNS = (
    "Record ID", "Original Baker Data", "Candidate Data", "Selected Lead ID", "Confidence", "Matching Method",
    "Evidence", "Conflict", "Recommended Action",
)
FAILED_COLUMNS = ("Record ID", "Original Baker Data", "Stage", "Error")
RUN_REPORT_COLUMNS = ("Metric", "Value")


class WorkbookValidationError(ValueError):
    """Raised when a workbook cannot be mapped to the Baker input contract."""


class InputColumnMapping:
    """Maps supported Baker logical fields to headers without relying on their order."""

    supported = {
        "first_name": "First Name", "last_name": "Last Name", "linkedin_url": "LinkedIn URL",
        "company_name": "Company Name", "job_title": "Job Title", "member_email": "Member email",
        "secondary_email": "Secondary Email", "primary_phone_number": "Primary Phone Number",
    }
    required = frozenset({"first_name", "last_name"})

    def __init__(self, headers: Sequence[Any]) -> None:
        normalized = {_header_key(header): str(header).strip() for header in headers if _header_key(header)}
        self.columns = {field: normalized.get(_header_key(label)) for field, label in self.supported.items()}
        missing = [self.supported[field] for field in self.required if self.columns[field] is None]
        if missing:
            raise WorkbookValidationError(f"Missing required Baker columns: {', '.join(sorted(missing))}.")


def read_baker_workbook(path: Path) -> list[BakerRecord]:
    """Read the active worksheet while preserving every original Baker column/value."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration as error:
            raise WorkbookValidationError("Workbook contains no header row.") from error
        mapping = InputColumnMapping(headers)
        header_names = [str(header).strip() if _header_key(header) else f"Column {index + 1}" for index, header in enumerate(headers)]
        records: list[BakerRecord] = []
        duplicate_counts: Counter[str] = Counter()
        for row_number, cells in enumerate(rows, start=2):
            original = {header_names[index]: _clean_value(value) for index, value in enumerate(cells)}
            if not any(value is not None for value in original.values()):
                continue
            fingerprint = hashlib.sha256(json.dumps(original, default=str, sort_keys=True).encode()).hexdigest()[:16]
            duplicate_counts[fingerprint] += 1
            records.append(BakerRecord(
                record_id=f"baker-{fingerprint}-{duplicate_counts[fingerprint]}", row_number=row_number,
                first_name=_string_value(original.get(mapping.columns["first_name"])),
                last_name=_string_value(original.get(mapping.columns["last_name"])),
                linkedin_url=_string_value(original.get(mapping.columns["linkedin_url"])),
                email=_first_present(original, mapping.columns["member_email"], mapping.columns["secondary_email"]),
                phone=_string_value(original.get(mapping.columns["primary_phone_number"])), original_values=original,
            ))
        return records
    finally:
        workbook.close()


def validate_baker_workbook(path: Path) -> InputColumnMapping:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        try:
            return InputColumnMapping(next(rows))
        except StopIteration as error:
            raise WorkbookValidationError("Workbook contains no header row.") from error
    finally:
        workbook.close()


class ExcelOutputWriter:
    """Writes all Phase 3 output schemas in a deterministic column order."""

    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir

    def write_enriched_contacts(self, contacts: Iterable[EnrichedContact]) -> Path:
        return self._write("enriched_contacts.xlsx", ENRICHED_COLUMNS, (_enriched_values(item) for item in contacts))

    def write_review_required(self, reviews: Iterable[ReviewRecord]) -> Path:
        return self._write("review_required.xlsx", REVIEW_COLUMNS, (_review_values(item) for item in reviews))

    def write_failed_records(self, failures: Iterable[FailedRecord]) -> Path:
        return self._write("failed_records.xlsx", FAILED_COLUMNS, (_failed_values(item) for item in failures))

    def write_run_report(self, report: RunReport) -> Path:
        return self._write("run_report.xlsx", RUN_REPORT_COLUMNS, ((key, value) for key, value in report.metrics.items()))

    def write_reverse_lookup(self, rows: Iterable[ReverseLookupRow], filename: str = "reverse_lookup.xlsx") -> Path:
        return self._write(filename, REVERSE_LOOKUP_COLUMNS, (_reverse_values(item) for item in rows))

    def _write(self, filename: str, columns: Sequence[str], rows: Iterable[Sequence[Any]]) -> Path:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        path = self.output_dir / filename
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Data"
        worksheet.append(list(columns))
        for row in rows:
            worksheet.append([_cell_value(value) for value in row])
        workbook.save(path)
        workbook.close()
        return path


def flatten_text_list(values: Any) -> str | None:
    if not isinstance(values, (list, tuple)):
        return None
    parts = [_string_value(value) for value in values]
    return "; ".join(value for value in parts if value) or None


def flatten_nested_records(value: Any) -> str | None:
    """Flatten common position/school response shapes into readable semicolon-separated text."""
    if isinstance(value, dict):
        value = value.get("values", value.get("elements", value))
    if not isinstance(value, (list, tuple)):
        return None
    rows: list[str] = []
    for item in value:
        if isinstance(item, str):
            rows.append(item)
        elif isinstance(item, dict):
            values = [_string_value(item.get(key)) for key in ("title", "companyName", "schoolName", "name", "degreeName", "fieldOfStudy")]
            dates = [_date_text(item.get(key)) for key in ("startDate", "endDate")]
            text = " — ".join(part for part in values if part)
            date_range = " to ".join(part for part in dates if part)
            rows.append(" | ".join(part for part in (text, date_range) if part))
    return "; ".join(row for row in rows if row) or None


def reverse_lookup_row(
    data: ReverseLookupData, *, lookup_email: str | None, current_role: str | None = None, current_company: str | None = None
) -> ReverseLookupRow:
    """Convert Reverse Lookup `data` into the fixed, flat spreadsheet representation."""
    return ReverseLookupRow(
        id=data.id, display_name=data.display_name, first_name=data.first_name, last_name=data.last_name,
        headline=data.headline, current_role=current_role, current_company=current_company, summary=data.summary,
        company=data.company_name, location=data.location, photo_url=data.photo_url, profile_url=data.profile_url,
        report_profile_url=data.report_profile_url, connection_count=data.connection_count,
        connections_obfuscated=data.is_connection_count_obfuscated,
        public_profile=_string_value(data.public_profile), lookup_email=lookup_email,
        skills=flatten_text_list(data.skills), locale_country=_string_value(data.locale.get("country")),
        locale_language=_string_value(data.locale.get("language")), education=flatten_nested_records(data.schools),
        work_history=flatten_nested_records(data.positions),
    )


def _enriched_values(item: EnrichedContact) -> tuple[Any, ...]:
    return tuple(asdict(item).values())


def _reverse_values(item: ReverseLookupRow) -> tuple[Any, ...]:
    return tuple(asdict(item).values())


def _review_values(item: ReviewRecord) -> tuple[Any, ...]:
    return (item.original_record.record_id, _json_text(item.original_record.original_values), _json_text(item.candidate_data), item.selected_lead_id,
            item.confidence, item.matching_method, item.evidence, item.conflict, item.recommended_action)


def _failed_values(item: FailedRecord) -> tuple[Any, ...]:
    return (item.original_record.record_id, _json_text(item.original_record.original_values), item.stage, item.error)


def _first_present(values: dict[str, Any], *headers: str | None) -> str | None:
    for header in headers:
        value = _string_value(values.get(header)) if header else None
        if value:
            return value
    return None


def _header_key(value: Any) -> str:
    return value.strip().casefold() if isinstance(value, str) else ""


def _clean_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, str):
        return value.strip() or None
    return value


def _string_value(value: Any) -> str | None:
    value = _clean_value(value)
    return value if isinstance(value, str) else str(value) if value is not None else None


def _cell_value(value: Any) -> Any:
    return _clean_value(value)


def _date_text(value: Any) -> str | None:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, dict):
        year, month = value.get("year"), value.get("month")
        if isinstance(year, int):
            return f"{year:04d}-{month:02d}" if isinstance(month, int) else str(year)
    return _string_value(value)


def _json_text(value: Any) -> str:
    return json.dumps(value, default=str, sort_keys=True)
