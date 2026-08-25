from __future__ import annotations

from math import nan
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from baker_enrichment.api.contracts import ReverseLookupData
from baker_enrichment.cli import main
from baker_enrichment.excel import (
    ENRICHED_COLUMNS,
    REVERSE_LOOKUP_COLUMNS,
    ExcelOutputWriter,
    WorkbookValidationError,
    flatten_nested_records,
    flatten_text_list,
    read_baker_workbook,
    reverse_lookup_row,
    validate_baker_workbook,
)
from baker_enrichment.models import BakerRecord, EnrichedContact, FailedRecord, ReviewRecord, RunReport


def make_workbook(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


def test_reads_normal_workbook_without_fixed_column_order(tmp_path: Path) -> None:
    path = tmp_path / "baker.xlsx"
    make_workbook(path, ["Job Title", "Last Name", "First Name", "LinkedIn URL", "Member email", "Primary Phone Number"], [["CEO", "Lovelace", "Ada", "https://linkedin.com/in/ada?trk=x", "ada@example.test", "123"]])
    record = read_baker_workbook(path)[0]
    assert (record.first_name, record.last_name, record.linkedin_url, record.email, record.phone) == ("Ada", "Lovelace", "https://linkedin.com/in/ada?trk=x", "ada@example.test", "123")
    assert record.original_values["Job Title"] == "CEO"


def test_optional_columns_are_not_required_and_blank_nan_values_are_none(tmp_path: Path) -> None:
    path = tmp_path / "baker.xlsx"
    make_workbook(path, ["First Name", "Last Name", "LinkedIn URL"], [["  Ada  ", " ", nan]])
    record = read_baker_workbook(path)[0]
    assert record.first_name == "Ada"
    assert record.last_name is None
    assert record.linkedin_url is None
    assert record.original_values["LinkedIn URL"] is None


def test_missing_required_input_columns_fails_validation(tmp_path: Path) -> None:
    path = tmp_path / "baker.xlsx"
    make_workbook(path, ["First Name", "Company Name"], [["Ada", "Analytical Engine"]])
    with pytest.raises(WorkbookValidationError, match="Last Name"):
        validate_baker_workbook(path)


def test_validate_command_never_needs_api_configuration(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "baker.xlsx"
    make_workbook(path, ["First Name", "Last Name"], [["Ada", "Lovelace"]])
    monkeypatch.delenv("ENRICH_API_KEY", raising=False)
    assert main(["validate", "--input", str(path)]) == 0


def test_duplicate_rows_are_preserved_with_unique_stable_ids(tmp_path: Path) -> None:
    path = tmp_path / "baker.xlsx"
    make_workbook(path, ["First Name", "Last Name"], [["Ada", "Lovelace"], ["Ada", "Lovelace"]])
    records = read_baker_workbook(path)
    assert len(records) == 2
    assert records[0].record_id.rsplit("-", 1)[0] == records[1].record_id.rsplit("-", 1)[0]
    assert records[0].record_id != records[1].record_id


def test_enriched_output_has_exact_schema_and_preserves_url(tmp_path: Path) -> None:
    path = ExcelOutputWriter(tmp_path / "nested" / "output").write_enriched_contacts([
        EnrichedContact(lead_id="lead-1", linkedin_profile="https://www.linkedin.com/in/ada/?trk=profile", email="ada@example.test")
    ])
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook.active
    assert tuple(cell.value for cell in worksheet[1]) == ENRICHED_COLUMNS
    assert worksheet.cell(2, 5).value == "https://www.linkedin.com/in/ada/?trk=profile"
    assert worksheet.cell(2, 26).value == "ada@example.test"
    workbook.close()


def test_reverse_lookup_flattening_is_readable_and_has_exact_schema(tmp_path: Path) -> None:
    row = reverse_lookup_row(
        ReverseLookupData(
            id="person-1", skills=("retail", "marketing strategy", "leadership"),
            locale={"country": "US", "language": "en"},
            schools={"values": [{"schoolName": "MIT", "degreeName": "BS", "startDate": {"year": 2000}, "endDate": {"year": 2004}}]},
            positions={"values": [{"title": "CEO", "companyName": "Example", "startDate": {"year": 2020, "month": 1}}]},
        ), lookup_email="ada@example.test",
    )
    path = ExcelOutputWriter(tmp_path).write_reverse_lookup([row])
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook.active
    assert tuple(cell.value for cell in worksheet[1]) == REVERSE_LOOKUP_COLUMNS
    assert worksheet.cell(2, 18).value == "retail; marketing strategy; leadership"
    assert worksheet.cell(2, 21).value == "MIT — BS | 2000 to 2004"
    assert worksheet.cell(2, 22).value == "CEO — Example | 2020-01"
    workbook.close()


def test_all_required_output_files_are_written_and_readable(tmp_path: Path) -> None:
    source = BakerRecord("baker-1", 2, "Ada", "Lovelace", None, None, None, {"First Name": "Ada"})
    writer = ExcelOutputWriter(tmp_path / "output")
    paths = [
        writer.write_enriched_contacts([]),
        writer.write_review_required([ReviewRecord(source, {"id": "lead-1"}, "lead-1", 0.5, "llm_fallback", "ambiguous", "conflict", "Review identity")]),
        writer.write_failed_records([FailedRecord(source, "search", "timeout")]),
        writer.write_run_report(RunReport({"Total records": 1, "Failures": 0})),
    ]
    assert [path.name for path in paths] == ["enriched_contacts.xlsx", "review_required.xlsx", "failed_records.xlsx", "run_report.xlsx"]
    assert all(path.is_file() and load_workbook(path).active.max_row >= 1 for path in paths)


def test_flatten_helpers_do_not_emit_json_syntax() -> None:
    assert flatten_text_list(["retail", "marketing"]) == "retail; marketing"
    assert flatten_nested_records([{"title": "Director", "companyName": "Example"}]) == "Director — Example"
